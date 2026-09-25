#!/usr/bin/env python3
"""
build_supplementary_table.py
----------------------------
Builds Supplementary Table 6b: every human phosphosite in the PDB survey that
is both buried and fully ionically coordinated within a single protein chain,
with a representative entry for each.

The table answers the part of Reviewer #2, major point 6 that asks whether the
coordinated buried phosphosites predicted here are seen in experimentally
determined structures.

Choice of representative entry
    Most sites appear in several entries. The one shown is the entry and chain
    with the most coordinating side chains, then the lowest accessibility, then
    the best resolution, among those where the site is both buried and fully
    coordinated.

Conservative reference check
    One column repeats the burial test with the reference accessibility of the
    unmodified parent residue in place of the derived phosphoresidue value.
    This is the more demanding of the two conventions, because the parent
    reference is smaller, and it identifies the sites whose burial does not
    depend on the derivation in phospho_reference_asa.py.

Legend and notes
    The legend is kept short; the definitions behind it are written as notes
    beneath the table. Every number in both is computed from the input tables,
    and the number of structures searched and the release cut-off are read from
    the header of coordinated_phosphoresidues.tsv, so that neither can go
    stale.

Input
-----
coordinated_phosphoresidues.tsv from find_coordinated_phosphoresidues.py,
phosphosites_classified.tsv from classify_kinase_sites.py, and the cached
UniProt records under pdb_cache/uniprot/ for gene and protein names.

Output
------
Supplementary_Table_S6b.tsv and Supplementary_Table_S6b.xlsx, with the notes
that define the table written beneath it (and as a commented header in the TSV),
and the matching paragraph of the legend with those notes in
Supplementary_Table_S6b_legend.txt. Given the curated literature table with
--table6a, also the complete two-tab Supplementary_Table_6.xlsx.

Usage
-----
    python build_supplementary_table.py
    python build_supplementary_table.py --table6a "Supplementary Table 6a.xlsx"

Dependencies
------------
    pip install pandas openpyxl
"""

import argparse
import json
import os
import re

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

RELASA_CUTOFF = 0.25  # buried at or below this (Levy 2010)
MIN_CONTACTS  = 2     # fully coordinated at or above this many Lys/Arg side chains
DIST_CUTOFF   = 4.0   # Å, ionic contact (Supplementary Table 4)

CONTACTS_FILE = "coordinated_phosphoresidues.tsv"
SITES_FILE    = "phosphosites_classified.tsv"
UNIPROT_CACHE = "pdb_cache/uniprot"
OUT_TSV       = "Supplementary_Table_S6b.tsv"
OUT_XLSX      = "Supplementary_Table_S6b.xlsx"
OUT_LEGEND    = "Supplementary_Table_S6b_legend.txt"
OUT_COMBINED  = "Supplementary_Table_6.xlsx"

PHOS_LABEL = {"SEP": "pSer", "TPO": "pThr", "PTR": "pTyr"}
NUMBER_WORDS = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five"}

# Tien et al. 2013 Empirical reference for the unmodified parent residue, used
# for the conservative burial check.
PARENT_MAX_ASA = {"SEP": 143.0, "TPO": 163.0, "PTR": 255.0}

EVIDENCE_LABEL = {
    "experimental":  "Experimental (ECO:0000269)",
    "by_similarity": "By similarity or large-scale (ECO:0000250, ECO:0007744, ECO:0000305)",
    "predicted":     "Predicted",
    "not_annotated": "Not annotated in UniProt",
}

REGION_LABEL = {
    "activation_loop_N":      "Activation loop, N-terminal segment (from the DFG motif)",
    "activation_loop_insert": "Activation loop, variable insert",
    "activation_loop_C":      "Activation loop, C-terminal segment (up to the APE motif)",
    "kinase_domain_other":    "Kinase domain, outside the activation loop",
    "outside_kinase_domain":  "Kinase, outside the kinase domain",
    "not_a_kinase":           "Not a protein kinase",
    "not_in_alignment":       "Kinase absent from the Kincore alignment",
}

# Paragraph b of the Supplementary Table 6 legend. The definitions behind it are
# kept short here and given in full in NOTES, which are written beneath the
# table itself. Every number is filled in from the data.
LEGEND = (
    "b, Systematic survey of the PDB. All {n_rows} human phosphosites that are both buried "
    "(RSA \u2264 {relasa_cutoff}, computed on the isolated chain) and held by two or more Lys or "
    "Arg side chains of the same chain, among all high-resolution human crystal structures "
    "containing phosphoserine, phosphothreonine or phosphotyrosine. Search criteria and "
    "definitions are given in the notes below the table.")

NOTES = [
    "PDB search: Every human X-ray structure in the PDB at \u2264 2.0 \u00c5 resolution, released up "
    "to and including {cutoff}, containing phosphoserine, phosphothreonine or phosphotyrosine "
    "({n_searched} entries). This gave {n_sites} unique phosphosites in {n_proteins} proteins with "
    "a phosphate coordinated by Lys or Arg, of which the {n_rows} listed are also buried in a chain "
    "in which they are fully coordinated (\u2265 {min_contacts} Arg/Lys contacts).",

    "Coordination: A Lys or Arg side chain of the same chain with a charged-group atom (Lys NZ; "
    "Arg NE, CZ, NH1, NH2) within {dist_cutoff} \u00c5 of a phosphate atom, the ionic-contact "
    "criterion of Supplementary Table 4. Fully coordinated: {min_contacts_word} or more such "
    "side chains.",

    "Burial: Relative solvent accessibility (RSA) \u2264 {relasa_cutoff}, computed with DSSP on "
    "the isolated chain, matching the monomeric analysis elsewhere in this work, and normalised "
    "to the empirical maxima of Tien et al. (2013). For phosphorylated residues, the empirical "
    "maximum of the parent residue was scaled by the phospho/parent ratio obtained by repeating "
    "the Gly-X-Gly enumeration of Tien et al.: SEP 226 \u00c5\u00b2, TPO 243 \u00c5\u00b2, "
    "PTR 334 \u00c5\u00b2. 'RSA using the unmodified parent reference' is the lowest RSA "
    "recalculated with the smaller maxima of unmodified Ser, Thr or Tyr. Entries in which a "
    "partner chain packs against the site show a lower accessibility for the whole entry than "
    "for the isolated chain, as expected.",

    "Exclusions: Co-crystallised substrate phosphopeptides (chains under 30 residues), "
    "non-human chains, and phosphoresidues that SIFTS does not map to UniProt, such as those in "
    "expression tags. Each residue was mapped to UniProt individually through SIFTS and accepted "
    "only where the canonical sequence carries the corresponding Ser, Thr or Tyr. Ubiquitin, "
    "encoded by four genes with an identical sequence, is counted once.",

    "Numbering: Phosphosites (e.g. pSer65) follow UniProt numbering. The 'PDB residue number' "
    "column and the coordinating residues are numbered as deposited in the representative PDB "
    "entry.",

    "Representative entry: The entry and chain with the most coordinating side chains, then "
    "the lowest RSA, then the best resolution. 'Lowest SASA' and 'Lowest RSA' are the lowest "
    "values among the entries in which the site is coordinated by Lys or Arg. 'Highest RSA' is "
    "the highest value among all entries in which the phosphorylated residue is resolved (RSA "
    "values are capped at 1).",

    "Reproducibility: Whether the site is buried and fully coordinated in all, at least half, "
    "or fewer than half of the entries in which the phosphorylated residue is resolved. 'PDB "
    "entries in which the site is resolved' lists these entries, with those in which the site is "
    "buried and fully coordinated first and marked with an asterisk.",

    "Kinase structural context: Kinase domains are from the Kincore alignment of human protein "
    "kinase domains (Modi & Dunbrack, 2019), in which the activation loop runs from the DFG to "
    "the APE motif.",

    "Code and complete results: https://github.com/langbnj/ptms/tree/main/pipeline/ptms_coordinated_phosphosites",
]

# Column widths, in characters.
WIDTHS = {"No.": 5, "Gene": 10, "UniProt": 9, "Protein": 34, "Phosphosite": 12,
          "Representative PDB entry": 11, "Chain": 6, "Resolution (Å)": 10,
          "PDB residue number": 11, "SASA in the isolated chain (Å²)": 12,
          "RSA in the isolated chain": 12, "Lowest SASA across entries (Å²)": 12,
          "Lowest RSA across entries": 12, "Highest RSA across entries": 12,
          "RSA using the unmodified parent reference": 13,
          "Buried under both reference conventions": 13, "Arg/Lys contacts": 10,
          "Coordinating residues (charged-group distance, Å)": 44,
          "Entries in which the site is resolved": 12, "Entries buried and fully coordinated": 13,
          "Reproducibility": 16, "PDB entries in which the site is resolved": 36,
          "UniProt modification evidence": 30, "Protein kinase": 9,
          "Structural context": 34}
# Characters per wrapped line where the default estimate is too generous (capitals and digits)
LINE_CHARS = {"PDB entries in which the site is resolved": 30}
THREE_DECIMALS = ["RSA in the isolated chain", "Lowest RSA across entries",
                  "Highest RSA across entries", "RSA using the unmodified parent reference"]

# Formatting matches the literature tab of Supplementary Table 6 (Table 6a).
FONT         = "Helvetica Neue"
FONT_SIZE    = 10
HEADER_FILL  = "165278"
HEADER_TEXT  = "FEFFFE"
GRID_COLOUR  = "C8C8C8"
SHEET_TITLE  = "Supp. Table 6b"
LINE_HEIGHT  = 13   # points per wrapped line at this font size
NOTE_COLUMNS = 14   # notes span the first this many columns


# ---------------------------------------------------------------------------
# Table
# ---------------------------------------------------------------------------

def read_provenance(path):
    """Structures searched, release cut-off and retrieval date, from the contact table header."""
    provenance = {}
    with open(path) as fh:
        for line in fh:
            if not line.startswith("#"):
                break
            match = re.match(r"#\s+(structures_searched|release_cutoff|retrieved)\s+(\S+)",
                             line)
            if match:
                provenance[match.group(1)] = match.group(2)
    return provenance


def uniprot_names(acc):
    """Primary gene name and recommended protein name from the cached UniProt record."""
    path = os.path.join(UNIPROT_CACHE, f"{acc}.json")
    if not os.path.exists(path):
        return "", ""
    with open(path) as fh:
        entry = json.load(fh)
    genes = entry.get("genes") or [{}]
    gene = genes[0].get("geneName", {}).get("value", "")
    protein = (entry.get("proteinDescription", {}).get("recommendedName", {})
               .get("fullName", {}).get("value", ""))
    return gene, protein


def representative(contacts, site):
    """The entry and chain shown for a site, and its contacts, best first."""
    rows = contacts[(contacts.acc == site.acc) & (contacts.site == site.site)]
    rows = rows[(rows.relasa_chain <= RELASA_CUTOFF) & (rows.n_contacts >= MIN_CONTACTS)]
    pick = rows.sort_values(["n_contacts", "relasa_chain", "resolution", "pdb_id", "chain"],
                            ascending=[False, True, True, True, True]).iloc[0]
    coordinating = rows[(rows.pdb_id == pick.pdb_id) & (rows.chain == pick.chain)] \
        .sort_values("dist_charged")
    return pick, coordinating


def resolved_entries(resolved, site):
    """The entries in which a site is resolved, coordinated or not, listed with those
    in which it is buried and fully coordinated first and marked with an asterisk,
    and its highest relative accessibility among them."""
    rows = resolved[(resolved.acc == site.acc) & (resolved.site == site.site)]
    qualifies = (rows.relasa_chain <= RELASA_CUTOFF) & (rows.n_contacts >= MIN_CONTACTS)
    per_entry = qualifies.groupby(rows.pdb_id).any()
    # Must agree with the counts from summarise_phosphosites.py
    assert len(per_entry) == site.entries_total, (site.acc, site.site)
    assert int(per_entry.sum()) == site.entries_buried_and_coordinated, (site.acc, site.site)
    listed = ([f"{entry}*" for entry in sorted(per_entry[per_entry].index)]
              + sorted(per_entry[~per_entry].index))
    return ", ".join(listed), float(rows.relasa_chain.max())


def build_table(contacts, sites, resolved):
    """One row per buried, fully coordinated site."""
    rows = []
    for site in sites[sites.buried_and_coordinated].itertuples():
        pick, coordinating = representative(contacts, site)
        entries, max_relasa = resolved_entries(resolved, site)
        parent_relasa = site.min_asa / PARENT_MAX_ASA[site.phos_res]
        gene, protein = uniprot_names(site.acc)
        rows.append({
            "Gene": gene,
            "UniProt": site.acc,
            "Protein": protein or site.protein_name,
            "Phosphosite": f"{PHOS_LABEL[site.phos_res]}{int(site.site)}",
            "Representative PDB entry": pick.pdb_id,
            "Chain": pick.chain,
            "Resolution (Å)": round(float(pick.resolution), 2),
            "PDB residue number": int(pick.phos_resnum),
            "SASA in the isolated chain (Å²)": int(pick.asa_chain),
            "RSA in the isolated chain": round(float(pick.relasa_chain), 3),
            "Lowest SASA across entries (Å²)": int(site.min_asa),
            "Lowest RSA across entries": round(float(site.min_relasa), 3),
            "Highest RSA across entries": round(max_relasa, 3),
            "RSA using the unmodified parent reference": round(float(parent_relasa), 3),
            "Buried under both reference conventions":
                "Yes" if parent_relasa <= RELASA_CUTOFF else "No",
            "Arg/Lys contacts": int(pick.n_contacts),
            "Coordinating residues (charged-group distance, Å)": "; ".join(
                f"{row.coord_res.capitalize()}{int(row.coord_resnum)} "
                f"{row.coord_atom}\u2013{row.phos_atom} ({row.dist_charged:.2f})"
                for row in coordinating.itertuples()),
            "Entries in which the site is resolved": int(site.entries_total),
            "Entries buried and fully coordinated": int(site.entries_buried_and_coordinated),
            "Reproducibility": ("all entries" if site.fraction_of_entries == 1
                                else "at least half of entries" if site.fraction_of_entries >= 0.5
                                else "fewer than half of entries"),
            "PDB entries in which the site is resolved": entries,
            "UniProt modification evidence": EVIDENCE_LABEL.get(site.evidence, site.evidence),
            "Protein kinase": "Yes" if site.is_kinase else "No",
            "Structural context": REGION_LABEL.get(site.region, site.region),
        })

    # Sorted by relative accessibility, which is comparable across residue types.
    # Ties are broken on accession and site so that the row order is reproducible.
    table = pd.DataFrame(rows).sort_values(
        ["Lowest RSA across entries", "UniProt", "Phosphosite"])
    table = table.reset_index(drop=True)
    table.insert(0, "No.", range(1, len(table) + 1))
    return table


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def wrapped_lines(value, width, per_line=None):
    """Lines a value occupies when wrapped in a column of the given width."""
    per_line = per_line or max(1, int(width * 1.1))
    return sum(max(1, -(-len(part) // per_line)) for part in str(value).split("\n"))


def write_sheet(sheet, table, notes=()):
    """The table on one worksheet, header in the first row and notes beneath it."""
    thin = Side(style="thin", color=GRID_COLOUR)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    columns = list(table.columns)

    for j, name in enumerate(columns, 1):
        cell = sheet.cell(row=1, column=j, value=name)
        cell.font = Font(name=FONT, size=FONT_SIZE, bold=True, color=HEADER_TEXT)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.border = box
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 52

    for i, row in enumerate(table.itertuples(index=False), start=2):
        for j, value in enumerate(row, 1):
            cell = sheet.cell(row=i, column=j, value=value)
            cell.font = Font(name=FONT, size=FONT_SIZE)
            cell.border = box
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    last_row = len(table) + 1

    for name in THREE_DECIMALS:
        j = columns.index(name) + 1
        for i in range(2, last_row + 1):
            sheet.cell(row=i, column=j).number_format = "0.000"
    j = columns.index("Resolution (Å)") + 1
    for i in range(2, last_row + 1):
        sheet.cell(row=i, column=j).number_format = "0.00"

    for j, name in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(j)].width = WIDTHS.get(name, 14)

    # Row heights, from the number of wrapped lines in the fullest cell, so that
    # nothing is cut off whether or not the spreadsheet program refits them.
    for i, row in enumerate(table.itertuples(index=False), start=2):
        lines = max(wrapped_lines(value, WIDTHS.get(name, 14), LINE_CHARS.get(name))
                    for name, value in zip(columns, row))
        sheet.row_dimensions[i].height = LINE_HEIGHT * lines + 4
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"

    # Notes beneath the table, each spanning the width of the first columns.
    if notes:
        span = min(len(columns), NOTE_COLUMNS)
        span_width = sum(WIDTHS.get(name, 14) for name in columns[:span])
        row = last_row + 2
        sheet.cell(row=row, column=1, value="Notes:").font = Font(name=FONT, size=FONT_SIZE,
                                                                  bold=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        for note in notes:
            row += 1
            # Each note starts with a label ending in a colon, set in bold.
            label, _, text = note.partition(": ")
            value = CellRichText([TextBlock(InlineFont(rFont=FONT, sz=FONT_SIZE, b=True), label + ":"),
                                  TextBlock(InlineFont(rFont=FONT, sz=FONT_SIZE), " " + text)])
            cell = sheet.cell(row=row, column=1, value=value)
            cell.font = Font(name=FONT, size=FONT_SIZE)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
            sheet.row_dimensions[row].height = LINE_HEIGHT * wrapped_lines(note, span_width) + 4


def write_workbook(table, path, notes=(), table6a=None):
    """Table 6b as a workbook of its own, or as the second tab of Table 6a.

    With table6a, the curated literature table is opened, its first sheet named
    Supp. Table 6a, and Supp. Table 6b added after it, giving Supplementary Table 6 as a
    single file.
    """
    if table6a:
        workbook = load_workbook(table6a)
        workbook.worksheets[0].title = "Supp. Table 6a"
        if SHEET_TITLE in workbook.sheetnames:
            del workbook[SHEET_TITLE]
        sheet = workbook.create_sheet(SHEET_TITLE)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_TITLE
    write_sheet(sheet, table, notes)
    workbook.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def format_date(iso):
    """2026-09-16 as 16 September 2026."""
    try:
        year, month, day = (int(x) for x in iso.split("-"))
    except ValueError:
        return iso
    months = ["January", "February", "March", "April", "May", "June", "July",
              "August", "September", "October", "November", "December"]
    return f"{day} {months[month - 1]} {year}"


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--contacts", default=CONTACTS_FILE,
                        help=f"contact table (default {CONTACTS_FILE})")
    parser.add_argument("--sites", default=SITES_FILE,
                        help=f"classified site table (default {SITES_FILE})")
    parser.add_argument("--tsv", default=OUT_TSV, help=f"output TSV (default {OUT_TSV})")
    parser.add_argument("--xlsx", default=OUT_XLSX, help=f"output workbook (default {OUT_XLSX})")
    parser.add_argument("--table6a",
                        help="the curated literature table; if given, Table 6b is added to it "
                             "as a second tab and saved as --combined")
    parser.add_argument("--combined", default=OUT_COMBINED,
                        help=f"two-tab Supplementary Table 6 (default {OUT_COMBINED})")
    args = parser.parse_args()

    contacts = pd.read_csv(args.contacts, sep="\t", comment="#")
    resolved = contacts[contacts["include"] == True]
    contacts = resolved[resolved.n_contacts >= 1]
    sites = pd.read_csv(args.sites, sep="\t", comment="#")
    provenance = read_provenance(args.contacts)

    table = build_table(contacts, sites, resolved)
    values = dict(
        n_rows=len(table),
        cutoff=format_date(provenance.get("release_cutoff", "?")),
        n_searched=(f"{int(provenance['structures_searched']):,}"
                    if "structures_searched" in provenance else "?"),
        n_sites=len(sites),
        n_proteins=sites.acc.nunique(),
        min_contacts=MIN_CONTACTS,
        min_contacts_word=NUMBER_WORDS.get(MIN_CONTACTS, str(MIN_CONTACTS)),
        dist_cutoff=DIST_CUTOFF,
        relasa_cutoff=RELASA_CUTOFF)
    legend = LEGEND.format(**values)
    notes = [note.format(**values) for note in NOTES]

    # The TSV carries the legend and notes as a commented header.
    with open(args.tsv, "w") as fh:
        for line in [legend, ""] + notes:
            fh.write(f"# {line}\n" if line else "#\n")
    table.to_csv(args.tsv, sep="\t", index=False, mode="a")
    write_workbook(table, args.xlsx, notes)
    with open(OUT_LEGEND, "w") as fh:
        fh.write(legend + "\n\nNotes:\n" + "\n".join(notes) + "\n")
    written = [args.tsv, args.xlsx, OUT_LEGEND]
    if args.table6a:
        write_workbook(table, args.combined, notes, table6a=args.table6a)
        written.append(args.combined)

    print(f"Supplementary Table 6b:\n >> {len(table)} sites, written to {', '.join(written)}\n")
    print(table[["No.", "Gene", "Phosphosite", "Representative PDB entry",
                 "Lowest SASA across entries (Å²)", "Lowest RSA across entries",
                 "Buried under both reference conventions", "Arg/Lys contacts",
                 "Reproducibility", "Structural context"]].to_string(index=False))
    print(f"\nLegend, paragraph b:\n{legend}\n\nNotes:\n" + "\n".join(notes))


if __name__ == "__main__":
    main()
